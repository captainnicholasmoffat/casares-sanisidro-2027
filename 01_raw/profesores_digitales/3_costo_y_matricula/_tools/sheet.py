import openpyxl,sys
def dump(fn,name='San Isidro'):
    wb=openpyxl.load_workbook(fn,read_only=True,data_only=True)
    target=None
    for ws in wb.worksheets:
        if ws.title in ('Carátula','Índice'): continue
        rows=list(ws.iter_rows(values_only=True,max_row=4))
        t=' '.join(str(c) for r in rows for c in r if c).replace('\xa0',' ')
        if name in t: target=ws; break
    out=[]
    for r in target.iter_rows(values_only=True):
        rr=[c for c in r if c not in (None,'')]
        if rr: out.append(' | '.join(str(c).replace('\xa0',' ').replace('\n',' ') for c in rr))
    open(fn.replace('.xlsx','_SanIsidro.txt'),'w').write(f'[hoja {target.title}]\n'+'\n'.join(out))
    return target.title,out
t,o=dump(sys.argv[1]); print('SHEET',t); print('\n'.join(o[:int(sys.argv[2]) if len(sys.argv)>2 else 200]))
