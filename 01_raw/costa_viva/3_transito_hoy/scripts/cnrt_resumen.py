# Extrae de los archivos oficiales de la CNRT (agosto 2026): pasajeros pagos por estación (Tren de la Costa y Mitre, estaciones de San Isidro)
# y trenes programados/corridos del Tren de la Costa y del ramal Tigre. Salida: cnrt_resumen.txt
import openpyxl, collections, warnings
warnings.filterwarnings('ignore')
B='/tmp/claude-0/-home-user-casares-sanisidro-2027/19687d82-9b19-5a6f-8ec4-f8d897955d9f/scratchpad/costa_viva_raw/v3_transito_hoy/trenes/'
out=open(B+'../cnrt_resumen.txt','w')
def p(*a): print(*a); print(*a,file=out)
for f,sh,est in [('cnrt_boletos/Boletos Tren de la Costa.xlsx','DATOS TDC#Bol por Estación',['Juan Anchorena','Las Barrancas','San Isidro R','Punta Chica','Maipú','Delta','Validadoras SUBE en Formación','Venta Sobre Tren','AJUSTE MESES ANTERIORES']),
                 ('cnrt_boletos/Boletos Mitre.xlsx','DATOS MIT#Bol por Estación',['Martínez','Acassuso','San Isidro','Beccar'])]:
    ws=openpyxl.load_workbook(B+f,read_only=True,data_only=True)[sh]
    agg=collections.defaultdict(lambda: collections.defaultdict(float)); meses=collections.defaultdict(set)
    for i,row in enumerate(ws.iter_rows(values_only=True)):
        if i==0 or row[0] is None: continue
        agg[row[0]][row[2]]+=row[3] or 0; meses[row[0]].add(row[1])
    p('##',f,'(pasajeros pagos por estación; año: meses con dato)')
    for y in range(2019,2027): p(y,len(meses[y]),{e:int(agg[y].get(e,0)) for e in est},'TOTAL archivo',int(sum(agg[y].values())))
wb=openpyxl.load_workbook(B+'ffcc_amba_cumplimiento_de_programa_2026-08_cnrt.xlsx',read_only=True,data_only=True)
p('## Cumplimiento de programa - Tren de la Costa (Año, Mes, Reg.Abs, Reg.Rel, Cumpl., Trenes prog. por día hábil, coches/tren, programados, cancelados, corridos, puntuales, atrasados, observaciones)')
for row in wb['TDC TAB'].iter_rows(values_only=True):
    if row[0] in (2024,2025,2026) and isinstance(row[1],str): p(row[:13])
p('## Cumplimiento de programa - Mitre ramal Retiro-Tigre')
for row in wb['MIT TAB'].iter_rows(values_only=True):
    if row[14] in (2025,2026) and isinstance(row[15],str): p(row[14:27])
